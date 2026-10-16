#!/usr/bin/env python3
"""Choose the fewest samples to remove so no pair is related up to a given degree.

Reads KING --related output (.kin and/or .kin0), builds a graph whose edges are
pairs that KING itself calls related up to DEGREE (its InfType column: Dup/MZ, PO, FS, 2nd, 3rd - this
is based on shared IBD segments and can disagree with the Kinship column) OR have
Kinship >= 2^-(degree+1.5), and finds a minimum vertex cover per
connected component: brute force up to 18 samples, an exact MILP (scipy >= 1.9) for larger clusters,
and only if scipy is missing / the MILP times out a greedy + redundancy-pruning fallback, in which case
a WARNING says the removal set is not guaranteed minimal.

Which sample goes when there is a choice (the minimum number of removals always comes first):
  1. a sample whose Population in the mega table is PREFER_REMOVE_POP (default "Iranian (ethnicity not
     reported)") is removed in preference to one with a known population;
  2. then the sample with the higher missing-genotype rate.
The table is given by the environment variable PREFER_REMOVE_TABLE (.xlsx sheet "mega_table", .csv or .tsv
with columns IID and Population); without it only rule 2 applies.

usage: prune_related.py DEGREE IMISS KIN [KIN ...] -- OUT_PREFIX
"""
import csv
import os
import re
import sys
from collections import defaultdict
from itertools import combinations

EXACT_MAX = 18          # components up to this size are solved by brute force (no scipy needed)
MILP_SECONDS = 120      # time limit per component for the exact (scipy MILP) solver

degree = int(sys.argv[1])
imiss = sys.argv[2]
sep = sys.argv.index("--")
kin_files = sys.argv[3:sep]
out = sys.argv[sep + 1]
thr = 2 ** (-(degree + 1.5))

fmiss = {}
with open(imiss) as fh:
    h = fh.readline().split()
    fi, ii, mi = h.index("FID"), h.index("IID"), h.index("F_MISS")
    for line in fh:
        p = line.split()
        if p:
            try:
                fmiss[(p[fi], p[ii])] = float(p[mi])
            except ValueError:
                fmiss[(p[fi], p[ii])] = 0.0

# samples to remove in preference to others (e.g. Population == "Iranian (ethnicity not reported)")
PREFER_POP = os.environ.get("PREFER_REMOVE_POP", "Iranian (ethnicity not reported)")


def load_prefer(path, pop):
    """IIDs whose Population equals `pop` in the mega table (xlsx sheet 'mega_table', csv or tsv)."""
    want = pop.strip().lower()
    try:
        if path.lower().endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb["mega_table"] if "mega_table" in wb.sheetnames else wb[wb.sheetnames[0]]
            rows = ws.iter_rows(values_only=True)
        else:
            fh = open(path, newline="")
            rows = csv.reader(fh, delimiter="\t" if path.lower().endswith((".tsv", ".txt")) else ",")
        header = [str(c).strip() if c is not None else "" for c in next(rows)]
        i_id, i_pop = header.index("IID"), header.index("Population")
        return {str(r[i_id]).strip() for r in rows
                if r and len(r) > max(i_id, i_pop) and r[i_pop] is not None and str(r[i_pop]).strip().lower() == want}
    except Exception as e:                          # missing openpyxl, wrong sheet/columns, unreadable file
        print("NOTE: could not read PREFER_REMOVE_TABLE %s (%s: %s); the Population preference is OFF" % (path, type(e).__name__, e))
        return set()


prefer_ids = set()
if os.environ.get("PREFER_REMOVE_TABLE"):
    prefer_ids = load_prefer(os.environ["PREFER_REMOVE_TABLE"], PREFER_POP)
    print("removal preference: %d sample(s) labelled %r in the table" % (len(prefer_ids), PREFER_POP))


def pref(n):
    return 1.0 if n[1] in prefer_ids else 0.0


# KING's InfType -> degree. Matching is tolerant of spelling/case ("Dup/MZ", "MZ twin", "MZTwin", "1st", ...)
def inf_degree(label):
    lab = re.sub(r"[^a-z0-9]", "", label.lower())
    if "dup" in lab or "mz" in lab or "twin" in lab:
        return 0
    if lab in ("po", "fs", "1st", "1stdegree"):
        return 1
    m = re.fullmatch(r"([2-9])(?:st|nd|rd|th)(?:degree)?", lab)
    if m:
        return int(m.group(1))
    return None            # "UN"/"unrelated"/unknown

unknown_labels = set()
kinship = {}
n_by_inftype = 0
for path in kin_files:
    try:
        fh = open(path)
    except FileNotFoundError:
        continue
    with fh:
        h = fh.readline().split()
        if not h:
            continue
        # --related tables (.kin/.kin0) have a Kinship column; --ibdseg tables (.seg) only PropIBD (kinship ~ PropIBD/2)
        k = h.index("Kinship") if "Kinship" in h else None
        pi = h.index("PropIBD") if "PropIBD" in h else None
        if k is None and pi is None:
            print("NOTE: %s has neither a Kinship nor a PropIBD column; skipped" % path)
            continue
        i1, i2 = h.index("ID1"), h.index("ID2")
        f1 = h.index("FID1") if "FID1" in h else h.index("FID")
        f2 = h.index("FID2") if "FID2" in h else h.index("FID")
        it = h.index("InfType") if "InfType" in h else None
        for line in fh:
            p = line.split()
            if not p:
                continue
            kv = float(p[k]) if k is not None else float(p[pi]) / 2
            by_inf = False
            if it is not None and len(p) > it:
                d_inf = inf_degree(p[it])
                by_inf = d_inf is not None and d_inf <= degree
                if d_inf is None and re.sub(r"[^a-z]", "", p[it].lower()) not in ("un", "unrelated", "other"):
                    unknown_labels.add(p[it])
            if kv >= thr or by_inf:
                if by_inf and kv < thr:
                    n_by_inftype += 1
                a, b = (p[f1], p[i1]), (p[f2], p[i2])
                key = frozenset((a, b))
                kinship[key] = max(kv, kinship.get(key, kv))        # same pair in two tables: keep the larger

if unknown_labels:
    print("NOTE: unrecognised InfType label(s) ignored: " + ", ".join(sorted(unknown_labels)))

adj = defaultdict(set)
for pair in kinship:
    a, b = tuple(pair)
    adj[a].add(b)
    adj[b].add(a)

seen, comps = set(), []
for n in sorted(adj):
    if n in seen:
        continue
    stack, comp = [n], []
    seen.add(n)
    while stack:
        x = stack.pop()
        comp.append(x)
        for y in adj[x]:
            if y not in seen:
                seen.add(y)
                stack.append(y)
    comps.append(sorted(comp))

def fm(n):
    return fmiss.get(n, 0.0)

def solve_exact(comp, edges):
    """Minimum vertex cover by MILP (scipy >= 1.9). Ties: preferred Population first, then more missing data.
    Returns (set, proven_optimal) or None if scipy's MILP is unavailable."""
    try:
        import numpy as np
        from scipy.optimize import Bounds, LinearConstraint, milp
        from scipy.sparse import lil_matrix
    except Exception:
        return None
    idx = {x: i for i, x in enumerate(comp)}
    n = len(comp)
    # cost 1 per removal, minus tiny bonuses: the Population preference (eps1) outweighs any missingness
    # bonus (eps2), and all bonuses together stay below 1, so one fewer removal always wins.
    eps1 = 0.5 / (n + 1)
    eps2 = eps1 / (2 * (n + 1))
    c = np.array([1.0 - eps1 * pref(x) - eps2 * fm(x) for x in comp])
    A = lil_matrix((len(edges), n))
    for k, (u, v) in enumerate(edges):
        A[k, idx[u]] = 1
        A[k, idx[v]] = 1
    res = milp(c, constraints=LinearConstraint(A.tocsr(), lb=1, ub=np.inf), integrality=np.ones(n),
               bounds=Bounds(0, 1), options={"time_limit": MILP_SECONDS})
    if res.x is None:
        return None
    chosen = {comp[i] for i in range(n) if res.x[i] > 0.5}
    if not all(u in chosen or v in chosen for u, v in edges):
        return None
    return chosen, res.status == 0


def solve_greedy(comp, edges):
    rem, left = set(), set(edges)
    while left:
        deg = defaultdict(int)
        for u, v in left:
            deg[u] += 1
            deg[v] += 1
        pick = max(deg, key=lambda x: (deg[x], pref(x), fm(x), x))
        rem.add(pick)
        left = {(u, v) for u, v in left if u != pick and v != pick}
    for x in sorted(rem, key=lambda x: (pref(x), fm(x), x)):   # drop redundant removals (least preferred first)
        if all(m in rem for m in adj[x]):
            rem.discard(x)
    return rem


removed = set()
not_proven = []
for comp in comps:
    edges = [(a, b) for a in comp for b in adj[a] if a < b]
    if len(comp) <= EXACT_MAX:
        for k in range(1, len(comp) + 1):
            covers = [set(s) for s in combinations(comp, k)
                      if all(u in s or v in s for u, v in edges)]
            if covers:
                removed |= max(covers, key=lambda s: (sum(pref(x) for x in s), sum(fm(x) for x in s), sorted(s)))
                break
    else:
        res = solve_exact(comp, edges)
        if res is not None:
            rem, optimal = res
            if not optimal:
                not_proven.append((len(comp), "MILP time limit reached; best solution found"))
        else:
            rem = solve_greedy(comp, edges)
            not_proven.append((len(comp), "scipy>=1.9 not available; greedy result"))
        removed |= rem
for size, why in not_proven:
    print("WARNING: related cluster of %d samples: removal set NOT guaranteed minimal (%s)" % (size, why),
          file=sys.stderr)

with open(out + ".removed_related.txt", "w") as fh:
    for n in sorted(removed):
        fh.write("%s %s\n" % n)
with open(out + ".removed_related_details.tsv", "w") as fh:
    fh.write("FID\tIID\tn_related\tmax_kinship\trelated_kept_samples\tpreferred_population\n")
    for n in sorted(removed):
        kept = sorted(m for m in adj[n] if m not in removed)
        mk = max(kinship[frozenset((n, m))] for m in adj[n])
        fh.write("%s\t%s\t%d\t%.4f\t%s\t%s\n" % (n[0], n[1], len(adj[n]), mk,
                                                 ",".join(m[1] for m in kept) or "-", "yes" if pref(n) else "-"))
if prefer_ids:
    print("removed %d sample(s), of which %d carry the preferred Population label" % (len(removed), sum(1 for x in removed if pref(x))))
with open(out + ".npairs", "w") as fh:
    fh.write("%d\n" % len(kinship))
print("related pairs (KING InfType up to degree %d, or kinship >= %.4f): %d (%d of them only by InfType); "
      "samples involved: %d; removed: %d"
      % (degree, thr, len(kinship), n_by_inftype, len(adj), len(removed)))
